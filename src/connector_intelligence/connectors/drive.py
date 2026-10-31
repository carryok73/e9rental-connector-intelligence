"""Read-only Google Drive API integration, text extraction, and classification."""
import csv
from io import BytesIO, StringIO
from pathlib import Path
import re
from typing import Any, BinaryIO

from docx import Document
from googleapiclient.http import MediaIoBaseDownload
from openpyxl import load_workbook
from pypdf import PdfReader

DRIVE_SCOPE = "https://www.googleapis.com/auth/drive.readonly"
DRIVE_SCOPES = [DRIVE_SCOPE]
GOOGLE_DOC_MIME = "application/vnd.google-apps.document"
GOOGLE_SHEET_MIME = "application/vnd.google-apps.spreadsheet"
GOOGLE_SLIDES_MIME = "application/vnd.google-apps.presentation"
SHEETS_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

DRIVE_CATEGORIES = ("sales", "marketing", "verification", "office", "accounts")


class DriveConnectorError(RuntimeError):
    """Safe, payload-free error from Drive setup or API operations."""


class UnsupportedDriveFile(DriveConnectorError):
    """File type has no configured, honest text extraction path."""


class InaccessibleDriveFile(DriveConnectorError):
    """The API reports that a file cannot be downloaded or exported."""


_CATEGORY_RULES: tuple[tuple[str, tuple[re.Pattern[str], ...]], ...] = (
    ("verification", tuple(re.compile(pattern, re.I) for pattern in (
        r"(?<![a-z0-9])kyc(?![a-z0-9])", r"(?<![a-z0-9])verifications?(?![a-z0-9])",
        r"(?<![a-z0-9])identit(?:y|ies)(?![a-z0-9])", r"(?<![a-z0-9])licenses?(?![a-z0-9])",
        r"proof[\s_-]+of[\s_-]+(?:address|identity)",
    ))),
    ("accounts", tuple(re.compile(pattern, re.I) for pattern in (
        r"(?<![a-z0-9])invoices?(?![a-z0-9])", r"(?<![a-z0-9])payments?(?![a-z0-9])",
        r"(?<![a-z0-9])accounts?(?![a-z0-9])", r"(?<![a-z0-9])accounting(?![a-z0-9])",
        r"(?<![a-z0-9])receipts?(?![a-z0-9])",
        r"(?<![a-z0-9])purchase[\s_-]*orders?(?![a-z0-9])",
    ))),
    ("sales", tuple(re.compile(pattern, re.I) for pattern in (
        r"(?<![a-z0-9])quotations?(?![a-z0-9])", r"(?<![a-z0-9])quotes?(?![a-z0-9])",
        r"(?<![a-z0-9])customers?(?![a-z0-9])", r"(?<![a-z0-9])rentals?(?![a-z0-9])",
        r"(?<![a-z0-9])sales(?![a-z0-9])", r"(?<![a-z0-9])leads?(?![a-z0-9])",
    ))),
    ("marketing", tuple(re.compile(pattern, re.I) for pattern in (
        r"(?<![a-z0-9])campaigns?(?![a-z0-9])", r"(?<![a-z0-9])brochures?(?![a-z0-9])",
        r"(?<![a-z0-9])marketing(?![a-z0-9])", r"(?<![a-z0-9])newsletters?(?![a-z0-9])",
        r"(?<![a-z0-9])advertisements?(?![a-z0-9])",
    ))),
    ("office", tuple(re.compile(pattern, re.I) for pattern in (
        r"(?<![a-z0-9])office(?![a-z0-9])", r"(?<![a-z0-9])polic(?:y|ies)(?![a-z0-9])",
        r"(?<![a-z0-9])procedures?(?![a-z0-9])", r"(?<![a-z0-9])meetings?(?![a-z0-9])",
        r"(?<![a-z0-9])operations(?![a-z0-9])",
    ))),
)


def _read_source(source: str | Path | bytes | BinaryIO, filename: str | None) -> tuple[bytes, str]:
    if isinstance(source, bytes):
        return source, filename or ""
    if hasattr(source, "read"):
        data = source.read()
        try:
            source.seek(0)
        except (AttributeError, OSError):
            pass
        inferred_name = filename or str(getattr(source, "name", ""))
        return data if isinstance(data, bytes) else str(data).encode("utf-8"), inferred_name
    path = Path(source)
    return path.read_bytes(), filename or path.name


def _join_located_segments(segments: list[tuple[str, dict[str, Any]]]) -> dict[str, Any]:
    """Join source segments with the extractor's newline separator and index their spans."""
    pieces: list[str] = []
    locations: list[dict[str, Any]] = []
    offset = 0
    for index, (segment, location) in enumerate(segments):
        if index:
            pieces.append("\n")
            offset += 1
        start = offset
        pieces.append(segment)
        offset += len(segment)
        locations.append({"start": start, "end": offset, "location": location})
    return {"text": "".join(pieces), "locations": locations}


def _line_locations(text: str, *, markdown: bool = False) -> list[dict[str, Any]]:
    """Index physical text lines, preserving their original line-ending characters."""
    lines = text.splitlines(keepends=True)
    locations: list[dict[str, Any]] = []
    offset = 0
    heading_stack: list[str] = []
    fence: tuple[str, int] | None = None
    heading_pattern = re.compile(r"^ {0,3}(#{1,6})(?:[ \t]+|$)(.*?)[ \t]*#*[ \t]*$")
    fence_pattern = re.compile(r"^ {0,3}(`{3,}|~{3,})")

    for line_number, line in enumerate(lines, start=1):
        line_body = line.rstrip("\r\n")
        section: str | None = None
        if markdown:
            fence_match = fence_pattern.match(line_body)
            if fence_match:
                marker = fence_match.group(1)
                if fence is None:
                    fence = (marker[0], len(marker))
                elif marker[0] == fence[0] and len(marker) >= fence[1]:
                    fence = None
            elif fence is None:
                heading = heading_pattern.match(line_body)
                if heading:
                    level = len(heading.group(1))
                    title = heading.group(2).strip()
                    heading_stack = heading_stack[:level - 1]
                    heading_stack.append(title)
            if heading_stack:
                section = " > ".join(heading_stack)

        location: dict[str, Any] = {"line": line_number}
        if section:
            location["section"] = section
        locations.append({
            "start": offset,
            "end": offset + len(line),
            "location": location,
        })
        offset += len(line)
    return locations


def _csv_record_locations(text: str) -> list[dict[str, Any]]:
    """Map logical CSV records to their raw character ranges and physical line ranges."""
    physical_lines = text.splitlines(keepends=True)
    line_offsets = [0]
    for line in physical_lines:
        line_offsets.append(line_offsets[-1] + len(line))

    locations: list[dict[str, Any]] = []
    previous_end_line = 0
    try:
        reader = csv.reader(StringIO(text, newline=""))
        for record_number, _record in enumerate(reader, start=1):
            end_line = reader.line_num
            start_line = previous_end_line + 1
            start_offset = line_offsets[min(start_line - 1, len(physical_lines))]
            end_offset = line_offsets[min(end_line, len(physical_lines))]
            locations.append({
                "start": start_offset,
                "end": end_offset,
                "location": {
                    "record": record_number,
                    "line_start": start_line,
                    "line_end": end_line,
                },
            })
            previous_end_line = end_line
    except csv.Error:
        # Text extraction remains available for malformed CSV. In that case, only
        # physical line locations can be reported reliably.
        locations = []
        offset = 0
        for line_number, line in enumerate(physical_lines, start=1):
            locations.append({
                "start": offset,
                "end": offset + len(line),
                "location": {"line_start": line_number, "line_end": line_number},
            })
            offset += len(line)
    return locations


def _extract_text_result(source: str | Path | bytes | BinaryIO,
                         filename: str | None = None, *,
                         include_locations: bool) -> dict[str, Any]:
    """Shared format extraction implementation with optional locator collection."""
    data, name = _read_source(source, filename)
    suffix = Path(name).suffix.lower()
    stream = BytesIO(data)

    if suffix == ".pdf":
        reader = PdfReader(stream)
        page_texts: list[str] = []
        segments: list[tuple[str, dict[str, Any]]] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            page_texts.append(text)
            if include_locations:
                segments.append((text, {"page": page_number}))
        if include_locations:
            return _join_located_segments(segments)
        return {"text": "\n".join(page_texts)}

    if suffix == ".docx":
        document = Document(stream)
        texts: list[str] = []
        segments = []
        heading_stack: list[str] = []
        for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text
            texts.append(text)
            if include_locations:
                style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
                heading = re.match(r"^Heading\s+(\d+)\b", style_name, re.I)
                if heading:
                    level = int(heading.group(1))
                    heading_stack = heading_stack[:level - 1]
                    heading_stack.append(text.strip())
                location: dict[str, Any] = {"paragraph": paragraph_number}
                if style_name:
                    location["style"] = style_name
                if heading_stack:
                    location["section"] = " > ".join(heading_stack)
                segments.append((text, location))

        # Match the existing extraction order: all top-level paragraphs, then
        # table cells in table/row/cell order. No original page numbers are inferred.
        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                for cell_number, cell in enumerate(row.cells, start=1):
                    text = cell.text
                    texts.append(text)
                    if include_locations:
                        segments.append((text, {
                            "table": table_number,
                            "row": row_number,
                            "cell": cell_number,
                        }))
        if include_locations:
            return _join_located_segments(segments)
        return {"text": "\n".join(texts)}

    if suffix == ".xlsx":
        book = load_workbook(stream, read_only=True, data_only=True)
        try:
            texts: list[str] = []
            segments = []
            for sheet in book.worksheets:
                header = f"Worksheet: {sheet.title}"
                texts.append(header)
                if include_locations:
                    segments.append((header, {
                        "worksheet": sheet.title,
                        "kind": "worksheet_header",
                    }))
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    row_text = " | ".join(str(cell) for cell in row if cell is not None)
                    texts.append(row_text)
                    if include_locations:
                        segments.append((row_text, {
                            "worksheet": sheet.title,
                            "row": row_number,
                        }))
            if include_locations:
                return _join_located_segments(segments)
            return {"text": "\n".join(texts)}
        finally:
            book.close()

    if suffix in {".txt", ".md", ".markdown"}:
        text = data.decode("utf-8-sig", errors="replace")
        if include_locations:
            return {
                "text": text,
                "locations": _line_locations(text, markdown=suffix in {".md", ".markdown"}),
            }
        return {"text": text}

    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        if include_locations:
            return {"text": text, "locations": _csv_record_locations(text)}
        return {"text": text}

    raise UnsupportedDriveFile(f"Unsupported document extension: {suffix or '(none)'}")


def extract_text_with_locations(source: str | Path | bytes | BinaryIO,
                                filename: str | None = None) -> dict[str, Any]:
    """Extract document text and source locators with offsets into the returned text.

    The returned ``start`` is inclusive and ``end`` is exclusive. All offsets count
    Python string characters in the exact ``text`` value returned here.
    """
    return _extract_text_result(source, filename, include_locations=True)


def extract_text(source: str | Path | bytes | BinaryIO, filename: str | None = None) -> str:
    """Extract text from a local path or in-memory Drive download/export."""
    return _extract_text_result(source, filename, include_locations=False)["text"]


def classify_document(name: str, text: str) -> tuple[str, str]:
    """Return one category and a concise rationale using a documented priority order.

    When a document matches multiple groups, verification outranks accounts, then
    sales, marketing, and office. A generic office assignment is the transparent
    fallback when no stronger signal appears.
    """
    sample = f"{name}\n{text[:20000]}"
    for category, patterns in _CATEGORY_RULES:
        match = next((found for pattern in patterns if (found := pattern.search(sample))), None)
        if match:
            term = match.group(0).strip()
            return category, f"Contains the {term!r} signal; assigned to {category}."
    return "office", "No sales, marketing, verification, or accounts signal was found; assigned to general office material."


def _execute(request: Any, operation: str) -> dict[str, Any]:
    try:
        result = request.execute()
    except Exception as exc:
        raise DriveConnectorError(f"Google Drive API request failed during {operation}.") from exc
    if not isinstance(result, dict):
        raise DriveConnectorError(f"Google Drive API returned an invalid response during {operation}.")
    return result


def list_drive_files(service: Any, page_size: int = 100) -> list[dict[str, Any]]:
    """List files accessible to the authorized user, excluding trashed files/folders."""
    if not 1 <= page_size <= 1000:
        raise ValueError("page_size must be between 1 and 1000")
    fields = (
        "nextPageToken,incompleteSearch,files(id,name,mimeType,modifiedTime,createdTime,"
        "webViewLink,description,size,md5Checksum,parents,driveId,capabilities(canDownload))"
    )
    files: list[dict[str, Any]] = []
    page_token: str | None = None
    seen_tokens: set[str] = set()
    while True:
        args: dict[str, Any] = {
            "q": "trashed = false and mimeType != 'application/vnd.google-apps.folder'",
            "pageSize": page_size,
            "orderBy": "name",
            "corpora": "user",
            "includeItemsFromAllDrives": True,
            "supportsAllDrives": True,
            "fields": fields,
        }
        if page_token:
            args["pageToken"] = page_token
        response = _execute(service.files().list(**args), "files.list")
        if response.get("incompleteSearch"):
            raise DriveConnectorError("Google Drive reported an incomplete file listing; sync stopped.")
        files.extend(item for item in response.get("files", []) or [] if item.get("id"))
        next_token = response.get("nextPageToken")
        if not next_token:
            break
        next_token = str(next_token)
        if next_token in seen_tokens:
            raise DriveConnectorError("Google Drive pagination returned a repeated page token.")
        seen_tokens.add(next_token)
        page_token = next_token
    return files


def _download_bytes(request: Any, operation: str) -> bytes:
    output = BytesIO()
    try:
        downloader = MediaIoBaseDownload(output, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
    except Exception as exc:
        raise DriveConnectorError(f"Google Drive download/export failed during {operation}.") from exc
    return output.getvalue()


def download_drive_file(service: Any, metadata: dict[str, Any]) -> tuple[bytes, str]:
    """Download a supported binary file or export a Google Doc/Sheet in memory."""
    file_id = str(metadata.get("id", ""))
    name = str(metadata.get("name", ""))
    mime_type = str(metadata.get("mimeType", ""))
    if not file_id:
        raise DriveConnectorError("Drive file metadata is missing its file ID.")
    capabilities = metadata.get("capabilities") or {}
    if capabilities.get("canDownload") is False:
        raise InaccessibleDriveFile("Drive reports that this file cannot be downloaded.")

    if mime_type == GOOGLE_SLIDES_MIME:
        raise UnsupportedDriveFile("Google Slides extraction is not supported; no content was processed.")
    if mime_type == GOOGLE_DOC_MIME:
        request = service.files().export_media(fileId=file_id, mimeType="text/plain")
        return _download_bytes(request, "files.export Google Docs"), f"{name}.txt"
    if mime_type == GOOGLE_SHEET_MIME:
        request = service.files().export_media(fileId=file_id, mimeType=SHEETS_XLSX_MIME)
        return _download_bytes(request, "files.export Google Sheets"), f"{name}.xlsx"

    suffix = Path(name).suffix.lower()
    supported_extensions = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md"}
    if suffix not in supported_extensions:
        raise UnsupportedDriveFile(f"Unsupported Drive MIME type or file extension: {mime_type or suffix or '(unknown)'}")
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    return _download_bytes(request, "files.get media"), name


def get_drive_service(settings: Any) -> Any:
    """Create Drive v3 OAuth client using only drive.readonly and its own token file."""
    try:
        from google.auth.transport.requests import Request
        from google.oauth2.credentials import Credentials
        from google_auth_oauthlib.flow import InstalledAppFlow
        from googleapiclient.discovery import build
    except ImportError as exc:
        raise DriveConnectorError("Google API dependencies are not installed in this environment.") from exc

    client_secrets = Path(settings.google_client_secrets_file)
    token_path = Path(settings.google_drive_token_file)
    try:
        credentials = Credentials.from_authorized_user_file(str(token_path), DRIVE_SCOPES) if token_path.exists() else None
        if credentials and credentials.scopes and set(credentials.scopes) != {DRIVE_SCOPE}:
            raise DriveConnectorError("The saved Drive token has scopes beyond drive.readonly; reauthorize with a readonly-only token.")
        if not credentials or not credentials.valid:
            if credentials and credentials.expired and credentials.refresh_token:
                credentials.refresh(Request())
            else:
                flow = InstalledAppFlow.from_client_secrets_file(str(client_secrets), DRIVE_SCOPES)
                credentials = flow.run_local_server(port=0)
            token_path.parent.mkdir(parents=True, exist_ok=True)
            token_path.write_text(credentials.to_json(), encoding="utf-8")
        if credentials.scopes and set(credentials.scopes) != {DRIVE_SCOPE}:
            raise DriveConnectorError("Google authorization did not grant only drive.readonly.")
        return build("drive", "v3", credentials=credentials, cache_discovery=False)
    except DriveConnectorError:
        raise
    except Exception as exc:
        raise DriveConnectorError("Google Drive OAuth setup failed. Check local OAuth configuration.") from exc


def authorized_drive_account(service: Any) -> str:
    response = _execute(service.about().get(fields="user(emailAddress)"), "about.get")
    email = str((response.get("user") or {}).get("emailAddress", "")).strip().casefold()
    if not email:
        raise DriveConnectorError("Google Drive profile did not return an account email.")
    return email


# Backward-compatible helper retained for existing project callers.
def google_client(service: str, version: str, client_secrets: str, token_file: str):
    """OAuth client helper; callers must pass a service-specific token path."""
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build

    if service not in {"gmail", "drive"}:
        raise ValueError("service must be 'gmail' or 'drive'")
    scopes = ["https://www.googleapis.com/auth/gmail.readonly"] if service == "gmail" else DRIVE_SCOPES
    token_path = Path(token_file)
    creds = Credentials.from_authorized_user_file(str(token_path), scopes) if token_path.exists() else None
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            creds = InstalledAppFlow.from_client_secrets_file(client_secrets, scopes).run_local_server(port=0)
        token_path.parent.mkdir(parents=True, exist_ok=True)
        token_path.write_text(creds.to_json(), encoding="utf-8")
    return build(service, version, credentials=creds, cache_discovery=False)
