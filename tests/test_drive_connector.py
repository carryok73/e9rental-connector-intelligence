from io import BytesIO
import json
from types import SimpleNamespace

from docx import Document
from openpyxl import Workbook
import pytest
from pypdf import PdfWriter

import src.connector_intelligence.connectors.drive as drive
from src.connector_intelligence.connectors.drive import (
    DriveConnectorError,
    GOOGLE_DOC_MIME,
    GOOGLE_SHEET_MIME,
    GOOGLE_SLIDES_MIME,
    DRIVE_SCOPE,
    SHEETS_XLSX_MIME,
    classify_document,
    download_drive_file,
    extract_text,
    get_drive_service,
    list_drive_files,
)
from src.connector_intelligence.connectors.drive_store import ingest_drive_file, sync_drive
from src.connector_intelligence.db import connect


class Request:
    def __init__(self, result=None, content=b"", error=None):
        self.result = result
        self.content = content
        self.error = error

    def execute(self):
        if self.error:
            raise self.error
        return self.result


class FakeDownloader:
    def __init__(self, output, request):
        self.output = output
        self.request = request
        self.done = False

    def next_chunk(self):
        if self.request.error:
            raise self.request.error
        if not self.done:
            self.output.write(self.request.content)
            self.done = True
        return None, self.done


class FakeFiles:
    def __init__(self, pages=None, blobs=None, exports=None, errors=None):
        self.pages = pages or [{"files": []}]
        self.blobs = blobs or {}
        self.exports = exports or {}
        self.errors = errors or {}
        self.list_calls = []
        self.media_calls = []
        self.export_calls = []

    def list(self, **kwargs):
        self.list_calls.append(kwargs)
        if "list" in self.errors:
            return Request(error=self.errors["list"])
        page = self.pages[1 if kwargs.get("pageToken") else 0]
        return Request(result=page)

    def get_media(self, **kwargs):
        self.media_calls.append(kwargs)
        return self._content_request(self.blobs.get(kwargs["fileId"], b""), "media")

    def export_media(self, **kwargs):
        self.export_calls.append(kwargs)
        file_id = kwargs["fileId"]
        return self._content_request(self.exports.get(file_id, b""), f"export:{file_id}")

    def _content_request(self, content, operation):
        return Request(content=content, error=self.errors.get(operation))


class FakeAbout:
    def __init__(self, email):
        self.email = email

    def get(self, **kwargs):
        return Request(result={"user": {"emailAddress": self.email}})


class FakeDriveService:
    def __init__(self, pages=None, blobs=None, exports=None, email="sandbox@example.test", errors=None):
        self.file_resource = FakeFiles(pages=pages, blobs=blobs, exports=exports, errors=errors)
        self.about_resource = FakeAbout(email)

    def files(self):
        return self.file_resource

    def about(self):
        return self.about_resource


@pytest.fixture(autouse=True)
def patch_downloader(monkeypatch):
    monkeypatch.setattr(drive, "MediaIoBaseDownload", FakeDownloader)


def make_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Quarterly accounts", "Invoice total"])
    sheet.append(["Q1", 125])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_file_list_paginates_and_scopes_to_authorized_user():
    service = FakeDriveService(pages=[
        {"files": [{"id": "f1", "name": "a.txt"}], "nextPageToken": "next"},
        {"files": [{"id": "f2", "name": "b.pdf"}]},
    ])
    files = list_drive_files(service, page_size=1)
    assert [file["id"] for file in files] == ["f1", "f2"]
    assert [call.get("pageToken") for call in service.file_resource.list_calls] == [None, "next"]
    assert all(call["corpora"] == "user" and call["q"].startswith("trashed = false")
               for call in service.file_resource.list_calls)
    assert all(call["supportsAllDrives"] is True for call in service.file_resource.list_calls)


def test_binary_file_download_uses_drive_api_media():
    service = FakeDriveService(blobs={"pdf1": b"plain test bytes"})
    content, name = download_drive_file(service, {"id": "pdf1", "name": "memo.txt", "mimeType": "text/plain"})
    assert content == b"plain test bytes"
    assert name == "memo.txt"
    assert service.file_resource.media_calls == [{"fileId": "pdf1", "supportsAllDrives": True}]


def test_google_doc_and_sheet_export_formats():
    service = FakeDriveService(exports={"doc1": b"Doc body", "sheet1": make_xlsx()})
    doc_bytes, doc_name = download_drive_file(service, {"id": "doc1", "name": "Terms", "mimeType": GOOGLE_DOC_MIME})
    sheet_bytes, sheet_name = download_drive_file(service, {"id": "sheet1", "name": "Accounts", "mimeType": GOOGLE_SHEET_MIME})
    assert doc_bytes == b"Doc body" and doc_name == "Terms.txt"
    assert extract_text(doc_bytes, doc_name) == "Doc body"
    assert "Invoice total" in extract_text(sheet_bytes, sheet_name)
    assert service.file_resource.export_calls == [
        {"fileId": "doc1", "mimeType": "text/plain"},
        {"fileId": "sheet1", "mimeType": SHEETS_XLSX_MIME},
    ]


def test_xlsx_extraction_includes_worksheet_names_and_keeps_cell_values():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Campaign Budget"
    sheet.append(["Campaign", "Planned Budget", "Actual Budget"])
    sheet.append(["Summer Promotion", 50000, 45000])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    text = extract_text(output.getvalue(), "campaign.xlsx")
    assert "Worksheet: Campaign Budget" in text
    assert "Campaign | Planned Budget | Actual Budget" in text
    assert "Summer Promotion | 50000 | 45000" in text


def test_text_extraction_supports_txt_markdown_csv_docx_and_xlsx():
    assert extract_text(b"Hello", "notes.txt") == "Hello"
    assert extract_text(b"# Heading", "notes.md") == "# Heading"
    assert extract_text(b"a,b\n1,2", "rows.csv") == "a,b\n1,2"
    document = Document()
    document.add_paragraph("Rental process")
    word_file = BytesIO()
    document.save(word_file)
    assert "Rental process" in extract_text(word_file.getvalue(), "process.docx")
    assert "Quarterly accounts" in extract_text(make_xlsx(), "accounts.xlsx")
    pdf_writer = PdfWriter()
    pdf_writer.add_blank_page(width=72, height=72)
    pdf_file = BytesIO()
    pdf_writer.write(pdf_file)
    assert extract_text(pdf_file.getvalue(), "blank.pdf") == ""


@pytest.mark.parametrize(("name", "text", "expected"), [
    ("proposal.txt", "Customer rental quote", "sales"),
    ("campaign.md", "Spring newsletter", "marketing"),
    ("driver-kyc.txt", "Identity verification", "verification"),
    ("handbook.txt", "Office meeting policy", "office"),
    ("invoice-2026.txt", "Payment receipt", "accounts"),
])
def test_every_drive_category_has_a_reason(name, text, expected):
    category, reason = classify_document(name, text)
    assert category == expected
    assert reason and "assigned" in reason


def test_ambiguous_classification_uses_documented_priority():
    category, _ = classify_document("KYC-invoice-quote.txt", "marketing campaign")
    assert category == "verification"


def test_sync_is_idempotent_and_exports_source_metadata(tmp_path):
    metadata = {
        "id": "f1", "name": "Office_policy.txt", "mimeType": "text/plain",
        "modifiedTime": "2026-10-20T10:00:00Z", "webViewLink": "https://drive.example.test/f1",
        "capabilities": {"canDownload": True},
    }
    service = FakeDriveService(pages=[{"files": [metadata]}], blobs={"f1": b"Office policy"})
    conn = connect(":memory:")
    first = sync_drive(conn, service, expected_account="sandbox@example.test", output_dir=tmp_path)
    second = sync_drive(conn, service, expected_account="sandbox@example.test", output_dir=tmp_path)
    assert first["extracted_files"] == second["extracted_files"] == 1
    assert conn.execute("SELECT COUNT(*) FROM records WHERE source='drive'").fetchone()[0] == 1
    row = conn.execute("SELECT source_id,occurred_at,category,metadata_json FROM records WHERE source='drive'").fetchone()
    metadata_saved = json.loads(row["metadata_json"])
    assert row["source_id"] == "f1" and row["category"] == "office"
    assert row["occurred_at"] == metadata["modifiedTime"]
    assert metadata_saved["web_view_link"] == metadata["webViewLink"]
    assert metadata_saved["category_reason"]
    exported = json.loads((tmp_path / "drive_output.json").read_text(encoding="utf-8"))
    assert len(exported) == 1 and exported[0]["metadata"]["mime_type"] == "text/plain"


def test_unsupported_slides_and_other_mime_are_not_marked_extracted():
    conn = connect(":memory:")
    result = ingest_drive_file(conn, FakeDriveService(), {"id": "slide", "name": "Deck", "mimeType": GOOGLE_SLIDES_MIME})
    assert result["status"] == "unsupported" and result["category"] is None
    other = ingest_drive_file(conn, FakeDriveService(), {"id": "zip", "name": "archive.zip", "mimeType": "application/zip"})
    assert other["status"] == "unsupported" and other["extraction_error"]


def test_inaccessible_file_is_recorded_without_download_attempt():
    conn = connect(":memory:")
    service = FakeDriveService()
    result = ingest_drive_file(conn, service, {
        "id": "locked", "name": "locked.pdf", "mimeType": "application/pdf",
        "capabilities": {"canDownload": False},
    })
    assert result["status"] == "inaccessible"
    assert result["category"] is None
    assert service.file_resource.media_calls == []


def test_api_errors_are_stored_safely_and_empty_file_is_reported():
    conn = connect(":memory:")
    errored_service = FakeDriveService(errors={"media": RuntimeError("private content in response")})
    failed = ingest_drive_file(conn, errored_service,
                               {"id": "error", "name": "x.txt", "mimeType": "text/plain"})
    assert failed["status"] == "api_error"
    assert "private content" not in failed["extraction_error"]
    empty = ingest_drive_file(conn, FakeDriveService(blobs={"empty": b""}),
                              {"id": "empty", "name": "blank.txt", "mimeType": "text/plain"})
    assert empty["status"] == "empty"
    assert empty["category"] is None


def test_drive_account_mismatch_prevents_file_listing():
    service = FakeDriveService(email="personal@example.test")
    with pytest.raises(DriveConnectorError, match="does not match"):
        sync_drive(connect(":memory:"), service, expected_account="sandbox@example.test")
    assert service.file_resource.list_calls == []


def test_oauth_requests_drive_readonly_only_and_uses_separate_token(monkeypatch, tmp_path):
    from google_auth_oauthlib.flow import InstalledAppFlow
    import googleapiclient.discovery

    requested = {}

    class FakeCredentials:
        valid = True
        expired = False
        refresh_token = None
        scopes = [DRIVE_SCOPE]

        def to_json(self):
            return '{"dummy":"local-test-token"}'

    class FakeFlow:
        def run_local_server(self, port):
            return FakeCredentials()

    def fake_flow(client_file, scopes):
        requested["client_file"] = client_file
        requested["scopes"] = list(scopes)
        return FakeFlow()

    monkeypatch.setattr(InstalledAppFlow, "from_client_secrets_file", fake_flow)
    monkeypatch.setattr(googleapiclient.discovery, "build", lambda *args, **kwargs: (args, kwargs))
    settings = SimpleNamespace(
        google_client_secrets_file=str(tmp_path / "client.json"),
        google_drive_token_file=str(tmp_path / "drive_token.json"),
    )
    client = get_drive_service(settings)
    assert requested["scopes"] == [DRIVE_SCOPE]
    assert client[0] == ("drive", "v3")
    assert (tmp_path / "drive_token.json").exists()


def test_api_list_errors_are_safe_and_not_suppressed():
    service = FakeDriveService(errors={"list": RuntimeError("sensitive API body")})
    with pytest.raises(DriveConnectorError, match="files.list") as error:
        list_drive_files(service)
    assert "sensitive API body" not in str(error.value)
