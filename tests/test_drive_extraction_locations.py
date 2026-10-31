from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook
import pytest
from pypdf import PdfWriter

from src.connector_intelligence.connectors.drive import (
    UnsupportedDriveFile,
    extract_text,
    extract_text_with_locations,
)


def _two_page_text_pdf(page_texts):
    def escaped(text):
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("ascii")

    page_count = len(page_texts)
    page_ids = list(range(3, 3 + page_count))
    content_ids = list(range(3 + page_count, 3 + 2 * page_count))
    font_id = 3 + 2 * page_count
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        ("<< /Type /Pages /Kids [" + " ".join(f"{item} 0 R" for item in page_ids)
         + f"] /Count {page_count} >>").encode(),
    ]
    for content_id in content_ids:
        objects.append((
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode())
    for text in page_texts:
        stream = b"BT /F1 12 Tf 50 750 Td (" + escaped(text) + b") Tj ET"
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    pdf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_number, value in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{object_number} 0 obj\n".encode())
        pdf.extend(value + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(pdf)


def _assert_location_offsets_align(result):
    text = result["text"]
    for item in result["locations"]:
        assert 0 <= item["start"] <= item["end"] <= len(text)
        assert isinstance(item["location"], dict)


def test_pdf_locations_identify_one_based_pages():
    data = _two_page_text_pdf(["Rental quotation on first page", "Invoice status on second page"])
    result = extract_text_with_locations(data, "sample.pdf")
    assert "Rental quotation on first page" in result["text"]
    assert "Invoice status on second page" in result["text"]
    assert [item["location"]["page"] for item in result["locations"]] == [1, 2]
    assert "Rental quotation" in result["text"][result["locations"][0]["start"]:result["locations"][0]["end"]]
    assert "Invoice status" in result["text"][result["locations"][1]["start"]:result["locations"][1]["end"]]
    _assert_location_offsets_align(result)


def test_blank_pdf_page_keeps_page_location_with_empty_text():
    data = _two_page_text_pdf(["", "Page two text"])
    result = extract_text_with_locations(data, "blank-page.pdf")
    assert result["text"].endswith("Page two text")
    assert result["locations"][0] == {"start": 0, "end": 0, "location": {"page": 1}}
    assert result["locations"][1]["location"] == {"page": 2}
    _assert_location_offsets_align(result)


def test_docx_locations_cover_paragraphs_heading_context_and_table_cells():
    document = Document()
    document.add_heading("Operations", level=1)
    document.add_paragraph("Meeting notes")
    document.add_heading("Decisions", level=2)
    document.add_paragraph("Four day rental approved")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Task"
    table.cell(1, 0).text = "Asha"
    table.cell(1, 1).text = "Update checklist"
    output = BytesIO()
    document.save(output)

    result = extract_text_with_locations(output.getvalue(), "notes.docx")
    assert result["text"] == extract_text(output.getvalue(), "notes.docx")
    assert any(item["location"].get("paragraph") == 4 and
               item["location"].get("section") == "Operations > Decisions"
               for item in result["locations"])
    assert {item["location"].get("cell") for item in result["locations"]
            if item["location"].get("table") == 1 and item["location"].get("row") == 2} == {1, 2}
    for item in result["locations"]:
        span = result["text"][item["start"]:item["end"]]
        if item["location"].get("paragraph") == 4:
            assert span == "Four day rental approved"
        if item["location"].get("table") == 1 and item["location"].get("row") == 2:
            assert span in {"Asha", "Update checklist"}
    _assert_location_offsets_align(result)
    assert all("page" not in item["location"] for item in result["locations"])


def test_xlsx_locations_identify_worksheet_and_one_based_rows():
    workbook = Workbook()
    first = workbook.active
    first.title = "Campaign Budget"
    first.append(["Campaign", "Planned Budget"])
    first.append(["Summer", 50000])
    first.append([None, None])
    first.append(["Winter", 25000])
    second = workbook.create_sheet("Invoices")
    second.append(["Invoice", "Status"])
    second.append(["SYN-1", "Pending"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    result = extract_text_with_locations(output.getvalue(), "multiple.xlsx")
    assert result["text"] == extract_text(output.getvalue(), "multiple.xlsx")
    assert "Worksheet: Campaign Budget" in result["text"]
    assert "Worksheet: Invoices" in result["text"]
    assert any(item["location"] == {"worksheet": "Campaign Budget", "row": 4}
               and result["text"][item["start"]:item["end"]] == "Winter | 25000"
               for item in result["locations"])
    assert any(item["location"] == {"worksheet": "Invoices", "row": 2}
               and result["text"][item["start"]:item["end"]] == "SYN-1 | Pending"
               for item in result["locations"])
    header = next(item for item in result["locations"]
                  if item["location"].get("kind") == "worksheet_header")
    assert result["text"][header["start"]:header["end"]] == "Worksheet: Campaign Budget"
    _assert_location_offsets_align(result)


def test_csv_locations_track_logical_records_and_multiline_physical_ranges():
    raw = (b"Name,Details,Amount\r\n"
           b"Mira,\"rental\r\nservices\",42\r\n"
           b"Noor,invoice,19\r\n")
    result = extract_text_with_locations(raw, "records.csv")
    assert result["text"] == raw.decode("utf-8-sig", errors="replace")
    assert result["locations"][0]["location"] == {
        "record": 1, "line_start": 1, "line_end": 1,
    }
    multiline = result["locations"][1]
    assert multiline["location"] == {"record": 2, "line_start": 2, "line_end": 3}
    assert result["text"][multiline["start"]:multiline["end"]] == 'Mira,"rental\r\nservices",42\r\n'
    assert result["locations"][2]["location"] == {
        "record": 3, "line_start": 4, "line_end": 4,
    }
    _assert_location_offsets_align(result)


def test_txt_and_markdown_locations_preserve_line_endings_and_sections():
    txt = b"first line\r\nsecond line\n"
    text_result = extract_text_with_locations(txt, "notes.txt")
    assert text_result["text"] == txt.decode("utf-8-sig", errors="replace")
    assert [item["location"]["line"] for item in text_result["locations"]] == [1, 2]
    assert text_result["text"][text_result["locations"][0]["start"]:text_result["locations"][0]["end"]] == "first line\r\n"

    markdown = b"# Campaign\ntext\n## Budget\nallocation\n```md\n# not a heading\n```\nend\n"
    md_result = extract_text_with_locations(markdown, "plan.md")
    assert md_result["text"] == markdown.decode("utf-8-sig", errors="replace")
    budget_line = next(item for item in md_result["locations"]
                       if md_result["text"][item["start"]:item["end"]] == "allocation\n")
    assert budget_line["location"] == {"line": 4, "section": "Campaign > Budget"}
    code_heading = next(item for item in md_result["locations"]
                        if "not a heading" in md_result["text"][item["start"]:item["end"]])
    assert code_heading["location"]["section"] == "Campaign > Budget"
    _assert_location_offsets_align(md_result)


def test_empty_supported_documents_keep_expected_empty_text_behavior():
    for filename in ("empty.txt", "empty.md", "empty.csv"):
        result = extract_text_with_locations(b"", filename)
        assert result == {"text": "", "locations": []}

    blank_pdf = _two_page_text_pdf([""])
    pdf_result = extract_text_with_locations(blank_pdf, "blank.pdf")
    assert pdf_result["text"] == ""
    assert pdf_result["locations"] == [{"start": 0, "end": 0, "location": {"page": 1}}]

    document = Document()
    word_file = BytesIO()
    document.save(word_file)
    docx_result = extract_text_with_locations(word_file.getvalue(), "empty.docx")
    assert docx_result["text"] == ""
    assert docx_result["locations"] == []

    workbook = Workbook()
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    xlsx_result = extract_text_with_locations(output.getvalue(), "empty.xlsx")
    assert xlsx_result["text"] == "Worksheet: Sheet"
    assert xlsx_result["locations"][0]["location"]["kind"] == "worksheet_header"


def test_unsupported_format_keeps_existing_exception_contract():
    with pytest.raises(UnsupportedDriveFile):
        extract_text_with_locations(b"data", "archive.bin")
    with pytest.raises(UnsupportedDriveFile):
        extract_text(b"data", "archive.bin")


def test_extract_text_remains_a_string_compatible_wrapper():
    examples = [
        (b"one\ntwo", "sample.txt"),
        (b"# Heading\nbody", "sample.md"),
        (b"one,two\n1,2", "sample.csv"),
        (_two_page_text_pdf(["Page one", "Page two"]), "sample.pdf"),
    ]
    for data, filename in examples:
        detailed = extract_text_with_locations(data, filename)
        plain = extract_text(data, filename)
        assert isinstance(plain, str)
        assert plain == detailed["text"]
        _assert_location_offsets_align(detailed)



def test_string_extractor_does_not_build_unused_location_index(monkeypatch):
    from src.connector_intelligence.connectors import drive

    def should_not_be_called(text, *, markdown=False):
        raise AssertionError("extract_text should not allocate detailed line locations")

    monkeypatch.setattr(drive, "_line_locations", should_not_be_called)
    assert extract_text(b"one\ntwo", "notes.txt") == "one\ntwo"
